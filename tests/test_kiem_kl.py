"""Test bộ kiem_kl trên dữ liệu TỰ DỰNG (không chứa số liệu dự án thật).

Mỗi ca mô phỏng một kiểu lỗi đã gặp thật ở đợt kiểm KL 25/09/2026, kèm ca đối chứng phải KHÔNG
bị báo. Nhóm 'CODEX R*' là ca hồi quy tái hiện đúng phản biện Codex 26/09 (phoi-hop/CODEX_PHAN_BIEN.md).
"""
import time

import ezdxf
import pytest
from openpyxl import Workbook

import kiem_kl.cad as cad
from kiem_kl.cad import dem_doi_tuong, hatch_giao_doi_tuong
from kiem_kl.quet_loi import CHUA, _BoTinh, _mau_cong_thuc, dinh_dang_ket_qua, quet_loi_sheet, quet_sheet_la
from kiem_kl.vni import bo_dau, giai_ma_vni, trong_nhu_vni


# ================================================================ VNI
@pytest.mark.parametrize("vni,unicode_", [
    (" + Tröø caùc hoá ga 1.0x2.0 chieám choã (10 hoá) ", " + Trừ các hố ga 1.0x2.0 chiếm chỗ (10 hố) "),
    ("BAÛNG KHOÁI LÖÔÏNG COÂNG TAÙC", "BẢNG KHỐI LƯỢNG CÔNG TÁC"),
    ("ÑÒA ÑIEÅM: PHÖÔØNG", "ĐỊA ĐIỂM: PHƯỜNG"),
    ("ÑAØ GIAÈNG", "ĐÀ GIẰNG"),
    ("ROÄNG", "RỘNG"),
    ("Toång soá hoá ga döôùi loøng ñöôøng", "Tổng số hố ga dưới lòng đường"),
    ("Chieàu daøi gôø boù væa heø", "Chiều dài gờ bó vỉa hè"),
])
def test_giai_ma_vni(vni, unicode_):
    assert giai_ma_vni(vni) == unicode_


@pytest.mark.parametrize("chuoi", ["Gia công và lắp đặt ván khuôn", "rõ ràng", "Hoá đơn", "TÍCH", "CHUNG", ""])
def test_unicode_that_giu_nguyen(chuoi):
    assert giai_ma_vni(chuoi) == chuoi


@pytest.mark.parametrize("vao,ra", [
    ("Tröø hố ga", "Trừ hố ga"),            # CODEX R5: ô lẫn VNI + Unicode
    ("Ñoå bê tông", "Đổ bê tông"),          # không được biến 'tông' thành 'tơng'
    ("Müller bê tông", "Müller bê tông"),   # tên nước ngoài giữ nguyên
    ("boàn", "boàn"),                       # mơ hồ, không có bằng chứng -> không đoán
    ("hoá", "hoá"),
])
def test_codex_r5_vni_theo_tung_tu(vao, ra):
    assert giai_ma_vni(vao) == ra


def test_vni_ep_che_do():
    assert giai_ma_vni("boàn", che_do="vni") == "bồn"
    assert giai_ma_vni("Tröø", che_do="unicode") == "Tröø"


def test_bo_dau_ca_vni_lan_unicode():
    assert bo_dau("Tröø caùc boàn caây") == "tru cac bon cay"
    assert bo_dau("Trừ các bồn cây") == "tru cac bon cay"
    assert not trong_nhu_vni("Đường")


# ================================================================ Excel
def _bang(dong):
    """Bảng KL kiểu mẫu: tiêu đề VNI 2 dòng, B TT, C diễn giải, E dài, F rộng, G cao, H diện tích,
    I số lượng, J KL riêng, K KL chung."""
    wb = Workbook()
    ws = wb.active
    ws.title = "KLCT"
    for o, v in {"B6": "TT", "C6": "HAÏNG MUÏC COÂNG TAÙC", "F6": "KÍCH THÖÔÙC", "H6": "DIEÄN",
                 "I6": "PHAÀN ", "J6": "KHOÁI LÖÔÏNG", "E7": "DAØI", "F7": "ROÄNG", "G7": "CAO",
                 "H7": "TÍCH", "I7": "G/ NHAU", "J7": "RIEÂNG", "K7": "CHUNG"}.items():
        ws[o] = v
    for o, v in dong.items():
        ws[o] = v
    return wb, ws


def _quet(ws, **kw):
    loi, cot, pv = quet_loi_sheet(ws, **kw)
    return loi


def _loai(loi):
    return {(x["loai"], x["o"].split("!")[1]) for x in loi}


def test_do_cot_tu_tieu_de_vni():
    _, ws = _bang({"B10": 1, "J10": 1})
    _, cot, _ = quet_loi_sheet(ws)
    assert cot["dien_giai"] == "C" and cot["so_luong"] == "I"
    assert cot["kl_rieng"] == "J" and cot["kl_chung"] == "K"
    assert cot["dai"] == "E" and cot["rong"] == "F" and cot["dien_tich"] == "H"


def test_he_so_bo_sot_va_dau_khoan_tru():
    wb, ws = _bang({
        "B10": 1, "C10": "Lát gạch", "K10": "=SUM(J11:J15)",
        "C11": "Diện tích lát", "J11": 500,
        "C12": " + Tröø caùc hoá ga 1.0x2.0 chieám choã (10 hoá)", "E12": 1.0, "F12": 2.0, "I12": 10, "J12": "=-E12*F12",
        "C13": " + Tröø hố thu 0.5x1.0 (4 hố)", "E13": 0.5, "F13": 1.0, "I13": 4, "J13": "=-E13*F13*I13",
        "C14": " + Tröø hố kỹ thuật 1.0x1.0 (1 hố)", "E14": 1.0, "F14": 1.0, "I14": 1, "J14": "=-E14*F14",
        "C15": " + Tröø caùc boàn caây chieám choã (20 caùi)", "E15": 1.0, "F15": 1.0, "I15": 20, "J15": "=E15*F15*I15",
    })
    loi = _quet(ws)
    l = _loai(loi)
    assert ("HE_SO_BO_SOT", "J12") in l
    assert ("DAU_KHOAN_TRU", "J15") in l
    assert ("HE_SO_BO_SOT", "J13") not in l
    assert ("HE_SO_BO_SOT", "J14") not in l
    x = next(x for x in loi if x["loai"] == "HE_SO_BO_SOT")
    assert x["hien"] == pytest.approx(-2.0) and x["de_xuat"] == pytest.approx(-20.0)


def test_khong_bao_nham_khi_so_luong_nhan_o_cot_chung():
    wb, ws = _bang({"B10": 1, "C10": "Đất hữu cơ bồn cây", "E10": 1.2, "F10": 1.2, "G10": 0.7,
                    "I10": 50, "J10": "=E10*F10*G10", "K10": "=I10*J10"})
    assert not any(x["loai"] == "HE_SO_BO_SOT" for x in _quet(ws))


@pytest.mark.parametrize("k10", ["=I10+J10", "=I10*100"])
def test_codex_r3_cot_chung_nhac_I_nhung_khong_nhan_J(k10):
    wb, ws = _bang({"B10": 1, "E10": 2, "F10": 3, "I10": 10, "J10": "=E10*F10", "K10": k10})
    assert ("HE_SO_BO_SOT", "J10") in _loai(_quet(ws))


def test_codex_r4_dau_xet_ket_qua_cuoi():
    # J dương, K = -J: kết quả cuối âm đúng -> không báo
    wb, ws = _bang({"B10": 1, "C10": "Trừ hố", "E10": 2, "F10": 3, "I10": 2, "J10": "=E10*F10*I10", "K10": "=-J10"})
    assert not any(x["loai"] == "DAU_KHOAN_TRU" for x in _quet(ws))


def test_nhan_lech_kich_thuoc_chi_trong_khoi_da_so_khop():
    d = {"B10": 1, "C10": "Đắp cát", "K10": "=SUM(J11:J14)"}
    for r, (nhan, e, f) in zip(range(11, 15), [("1.0x2.0", 1.0, 2.0), ("1.5x1.5", 1.5, 1.5),
                                                 ("2.0x3.0", 2.0, 3.0), ("1.0x3.0", 1.5, 2.0)]):
        d.update({f"C{r}": f" + Trừ hố {nhan} (1 hố)", f"E{r}": e, f"F{r}": f, f"I{r}": 1, f"J{r}": f"=-E{r}*F{r}*I{r}"})
    d.update({"B20": 2, "C20": "Nắp đan", "K20": "=SUM(J21:J23)"})
    for r in (21, 22, 23):
        d.update({f"C{r}": " + Hố ga 1.2x1.6 (nắp 0.7x1.0)", f"E{r}": 0.7, f"F{r}": 1.0, f"J{r}": f"=E{r}*F{r}"})
    wb, ws = _bang(d)
    assert {x["o"].split("!")[1] for x in _quet(ws) if x["loai"] == "NHAN_LECH_KT"} == {"E14"}


def test_codex_r4_nhan_cat_khoi_tai_dau_muc():
    d = {"B10": 1, "B13": 2}
    for r in (11, 12):
        d.update({f"C{r}": " + hố 1x2", f"E{r}": 1, f"F{r}": 2, f"J{r}": f"=E{r}*F{r}"})
    d.update({"C14": " + hố 1x2", "E14": 3, "F14": 2, "J14": "=E14*F14"})
    wb, ws = _bang(d)
    assert not any(x["loai"] == "NHAN_LECH_KT" for x in _quet(ws))


def test_codex_r4_khac_phep_toan():
    d = {"B10": 1}
    for r in (11, 12):
        d.update({f"E{r}": 2, f"F{r}": 3, f"I{r}": 4, f"J{r}": f"=E{r}*F{r}*I{r}"})
    d.update({"E13": 2, "F13": 3, "I13": 4, "J13": "=E13*F13/I13"})
    wb, ws = _bang(d)
    assert ("CONG_THUC_LECH", "J13") in _loai(_quet(ws))


def test_codex_r4_giu_tham_chieu_tuyet_doi():
    assert _mau_cong_thuc("=E11*$I$1", 11) == _mau_cong_thuc("=E12*$I$1", 12) == "=E[0]*$I$1"


def test_so_luong_trong_dien_giai_khac_cot_so_luong():
    wb, ws = _bang({"B10": 1, "K10": "=SUM(J11:J11)",
                    "C11": " + Trừ hố ga 1.0x1.0 (12 hố)", "E11": 1.0, "F11": 1.0, "I11": 10, "J11": "=-E11*F11*I11"})
    assert ("SO_LUONG_LECH", "I11") in _loai(_quet(ws))


def test_sum_bo_sot_dong_cuoi():
    wb, ws = _bang({"B10": 1, "C10": "Hạng mục", "K10": "=SUM(J11:J12)",
                    "J11": 5, "J12": 6, "C13": " + dòng con bị sót", "J13": 7,
                    "B14": 2, "C14": "Hạng mục sau", "K14": "=SUM(J15:J15)", "J15": 1, "B16": 3})
    l = _loai(_quet(ws))
    assert ("SUM_BO_SOT", "K10") in l
    assert ("SUM_BO_SOT", "K14") not in l


def test_ref_va_lien_ket_ngoai_va_hang_so():
    wb, ws = _bang({"B10": 1, "K10": "=#REF!*2", "B11": 2, "K11": "=K10-'[old.xls]KLCT'!$K$67",
                    "B12": 3, "K12": "=180.434+23.56", "B13": 4, "E13": "=#REF!", "J13": "=E13*2"})
    l = _loai(_quet(ws))
    assert ("THAM_CHIEU", "K10") in l and ("THAM_CHIEU", "K11") in l
    assert ("HANG_SO", "K12") in l
    assert ("THAM_CHIEU", "E13") in l      # CODEX R4: #REF! ở cột kích thước


def test_so_go_tay_trong_cot_khoi_luong():
    wb, ws = _bang({"B10": 1, "K10": 123.45})
    assert ("HANG_SO", "K10") in _loai(_quet(ws))


def test_cong_thuc_lech_bo_qua_cot_trong():
    d = {"B10": 1, "K10": "=SUM(J11:J15)"}
    for r in (11, 12, 13, 14):
        d.update({f"E{r}": 1.0, f"F{r}": 2.0, f"G{r}": 0.5, f"I{r}": 2, f"J{r}": f"=-E{r}*F{r}*I{r}*G{r}"})
    d.update({"E15": 3.0, "F15": 0.2, "G15": 0.5, "J15": "=-E15*F15*G15"})
    wb, ws = _bang(d)
    assert not any(x["loai"] == "CONG_THUC_LECH" for x in _quet(ws))


def test_sheet_la_va_dinh_dang():
    wb, ws = _bang({"B10": 1, "K10": "=Phu!A1"})
    wb.create_sheet("Phu")
    wb.create_sheet("xxxxxxxx")
    wb.create_sheet("Cu")
    loi, khong_goi = quet_sheet_la(wb)
    assert [x["o"] for x in loi] == ["xxxxxxxx"]
    assert "Cu" in khong_goi and "Phu" not in khong_goi
    assert dinh_dang_ket_qua(loi, max_ket_qua=5).startswith("Tổng 1 nghi vấn: CAO 1")


# ---------------------------------------------------------------- CODEX R1/R2: bộ tính
def _ws_trong():
    wb = Workbook()
    return wb.active


@pytest.mark.parametrize("ct,ky_vong", [
    ("=1E3+2", 1002.0),
    ("=(2+3)*4-1", 19.0),
    ("=-2*-3", 6.0),
])
def test_codex_r1_bo_tinh_so(ct, ky_vong):
    assert _BoTinh(_ws_trong())._tinh(ct) == pytest.approx(ky_vong)


def test_codex_r1_round_va_sum_lan_truyen():
    ws = _ws_trong()
    ws["A1"], ws["A2"], ws["A3"] = "=ROUND(2.4,0)", "=A1+5", "=SUM(A1:A2)"
    bt = _BoTinh(ws)
    assert (bt.lay("A", 1), bt.lay("A", 2), bt.lay("A", 3)) == (2.0, 7.0, 9.0)


def test_codex_r1_ham_chua_ho_tro_lan_truyen_CHUA():
    ws = _ws_trong()
    ws["A1"], ws["A2"], ws["A3"] = "=VLOOKUP(1,B1:C2,2,0)", "=A1+5", "=SUM(A1:A2)"
    bt = _BoTinh(ws)
    assert bt.lay("A", 2) is CHUA and bt.lay("A", 3) is CHUA


@pytest.mark.parametrize("ct", ["=9**9", "=5//2", "=2^10", "=1E300*1E300", "=\"a\"&\"b\""])
def test_codex_r2_cu_phap_ngoai_ho_tro(ct):
    assert _BoTinh(_ws_trong())._tinh(ct) is CHUA


def test_codex_r2_dai_khong_lo_chan_truoc_khi_bung():
    t = time.monotonic()
    assert _BoTinh(_ws_trong())._tinh("=SUM(A1:XFD1048576)") is CHUA
    assert time.monotonic() - t < 1.0


def test_vong_tham_chieu():
    ws = _ws_trong()
    ws["A1"], ws["A2"] = "=A2+1", "=A1+1"
    assert _BoTinh(ws).lay("A", 1) is CHUA


# ---------------------------------------------------------------- CODEX R9: phạm vi
def test_codex_r9_max_dong_dung_va_bao_chua_quet():
    d = {"B10": 1}
    for r in range(10, 20):
        d[f"J{r}"] = r
    wb, ws = _bang(d)
    loi, _, pv = quet_loi_sheet(ws, dong_bat_dau=10, max_dong=1)
    assert (pv["dau"], pv["cuoi"]) == (10, 10)
    assert "CHƯA QUÉT dòng 11–19" in dinh_dang_ket_qua(loi, pham_vi=[pv])


def test_codex_r9_khong_co_cot_ket_qua_thi_bao_cau_truc():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["B2"] = "Diễn giải", "Khối lượng", "=#REF!"
    loi, _, _ = quet_loi_sheet(ws)
    assert loi and loi[0]["loai"] == "CAU_TRUC" and loi[0]["muc"] == "CAO"


def test_cot_map_truyen_tay_thi_khong_do():
    wb, ws = _bang({"B10": 1, "E10": 2, "F10": 3, "I10": 10, "J10": "=E10*F10"})
    loi, cot, _ = quet_loi_sheet(ws, cot_map={"so_luong": "I", "kl_rieng": "J", "dai": "E", "rong": "F"}, dong_bat_dau=10)
    assert cot == {"so_luong": "I", "kl_rieng": "J", "dai": "E", "rong": "F"}
    assert ("HE_SO_BO_SOT", "J10") in _loai(loi)


# ================================================================ DXF
def _hcn(x, y, w, h):
    return [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]


@pytest.fixture
def dxf_mau(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 6
    m = d.modelspace()
    b = d.blocks.new("HOGA")
    b.add_lwpolyline(_hcn(0, 0, 1.2, 1.6), close=True)
    b.add_attdef("MA", (0, 0), dxfattribs={"height": 5})  # chữ to: không được làm phình hình
    h = m.add_hatch(dxfattribs={"layer": "LAT"})
    h.paths.add_polyline_path(_hcn(0, 0, 20, 4), is_closed=True, flags=1)
    h.paths.add_polyline_path(_hcn(10, 1, 1, 1), is_closed=True, flags=16)
    for x, y in [(10, 1), (3, 1), (30, 1)]:
        m.add_lwpolyline(_hcn(x, y, 1, 1), close=True, dxfattribs={"layer": "BONCAY"})
    m.add_blockref("HOGA", (5, 1), dxfattribs={"layer": "HG"})
    m.add_blockref("HOGA", (19.7, 1), dxfattribs={"layer": "HG"})
    m.add_line((0, -1), (10, -1), dxfattribs={"layer": "BOVIA"})
    m.add_arc((0, 0), 10, 0, 90, dxfattribs={"layer": "BOVIA"})
    m.add_text("HG 12", dxfattribs={"layer": "TEXT"})
    p = tmp_path / "mau.dxf"
    d.saveas(p)
    return str(p)


def test_dem_doi_tuong(dxf_mau):
    txt = dem_doi_tuong(dxf_mau)
    assert "I|HOGA|HG|2" in txt
    assert "R|1.00x1.00|BONCAY|3" in txt
    assert "L|BOVIA|25.71" in txt
    assert "H|LAT|79.00" in txt
    assert "T|HG12|1" in txt


def test_hatch_giao_doi_tuong(dxf_mau):
    bon = hatch_giao_doi_tuong(dxf_mau, "LAT", "BONCAY")
    assert "TRONG_VUNG 1 · DA_KHOET 1 · MOT_PHAN 0 · NGOAI_VUNG 1" in bon
    assert "TỔNG DIỆN TÍCH CẦN TRỪ khỏi vùng lát = 1.000" in bon
    hg = hatch_giao_doi_tuong(dxf_mau, "LAT", "HOGA")
    assert "TRONG_VUNG 1 · DA_KHOET 0 · MOT_PHAN 1" in hg
    # hố ở mép: 0,3 m bề ngang nằm trong vùng (0,3×1,6 = 0,48) -> tổng trừ 1,92 + 0,48
    assert "TỔNG DIỆN TÍCH CẦN TRỪ khỏi vùng lát = 2.400" in hg


def test_hatch_khong_thay_layer(dxf_mau):
    assert "Không thấy HATCH" in hatch_giao_doi_tuong(dxf_mau, "KHONGCO", "BONCAY")


def test_codex_r6_don_vi_mm(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 4
    h = d.modelspace().add_hatch(dxfattribs={"layer": "LAT"})
    h.paths.add_polyline_path(_hcn(0, 0, 1000, 1000), is_closed=True, flags=1)
    d.modelspace().add_line((0, 0), (2000, 0), dxfattribs={"layer": "BV"})
    p = tmp_path / "mm.dxf"
    d.saveas(p)
    txt = dem_doi_tuong(str(p))
    assert "H|LAT|1.00" in txt and "L|BV|2.00" in txt


def test_codex_r6_don_vi_chua_ro(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 0
    d.modelspace().add_line((0, 0), (5, 0))
    p = tmp_path / "u0.dxf"
    d.saveas(p)
    assert "chưa xác định" in dem_doi_tuong(str(p))
    assert "L|0|5000.00" in dem_doi_tuong(str(p), he_so_don_vi=1000)


def test_codex_r7_lo_khoet_mot_phan_va_dao_nho(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 6
    m = d.modelspace()
    h = m.add_hatch(dxfattribs={"layer": "LAT"})
    h.paths.add_polyline_path(_hcn(0, 0, 10, 10), is_closed=True, flags=1)
    h.paths.add_polyline_path(_hcn(1, 1, 0.6, 1), is_closed=True, flags=16)
    h.paths.add_polyline_path(_hcn(20, 0, 4, 4), is_closed=True, flags=1)
    h.paths.add_polyline_path(_hcn(21, 1, 1, 1), is_closed=True, flags=16)
    m.add_lwpolyline(_hcn(1, 1, 1, 1), close=True, dxfattribs={"layer": "OBJ"})   # 60% lỗ, 40% vùng tô
    m.add_lwpolyline(_hcn(21, 1, 1, 1), close=True, dxfattribs={"layer": "OBJ"})  # trùng lỗ của đảo nhỏ
    p = tmp_path / "r7.dxf"
    d.saveas(p)
    txt = hatch_giao_doi_tuong(str(p), "LAT", "OBJ")
    assert "TRONG_VUNG 0 · DA_KHOET 1 · MOT_PHAN 1 · NGOAI_VUNG 0" in txt
    assert "TỔNG DIỆN TÍCH CẦN TRỪ khỏi vùng lát = 0.400" in txt


def test_codex_r8_block_xoay_do_hinh_that(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 6
    b = d.blocks.new("HCN")
    b.add_lwpolyline(_hcn(0, 0, 2, 1), close=True)
    m = d.modelspace()
    h = m.add_hatch(dxfattribs={"layer": "LAT"})
    h.paths.add_polyline_path(_hcn(-10, -10, 20, 20), is_closed=True, flags=1)
    m.add_blockref("HCN", (0, 0), dxfattribs={"rotation": 45, "layer": "OBJ"})
    p = tmp_path / "r8.dxf"
    d.saveas(p)
    txt = hatch_giao_doi_tuong(str(p), "LAT", "HCN")
    assert "TỔNG DIỆN TÍCH CẦN TRỪ khỏi vùng lát = 2.000" in txt


def test_codex_r8_bulge_khong_phai_chu_nhat(tmp_path):
    d = ezdxf.new()
    d.header["$INSUNITS"] = 6
    d.modelspace().add_lwpolyline([(0, 0, 0), (1, 0, 0.5), (1, 1, 0), (0, 1, 0)], format="xyb", close=True)
    p = tmp_path / "bulge.dxf"
    d.saveas(p)
    assert "R|" not in dem_doi_tuong(str(p))


def test_codex_r8_thieu_shapely_khong_xap_xi(dxf_mau, monkeypatch):
    monkeypatch.setattr(cad, "HAS_SHAPELY", False)
    txt = dem_doi_tuong(dxf_mau)
    assert "KHÔNG TÍNH (thiếu shapely" in txt and "H|LAT|" not in txt


@pytest.mark.parametrize("kw,chu", [({"dung_sai": 0}, "ngoài miền"), ({"layer_hatch": "("}, "không phải regex")])
def test_codex_r10_kiem_dau_vao(dxf_mau, kw, chu):
    tham = dict(layer_hatch="LAT", loc_doi_tuong="BONCAY")
    tham.update(kw)
    with pytest.raises(RuntimeError, match=chu):
        hatch_giao_doi_tuong(dxf_mau, **tham)
