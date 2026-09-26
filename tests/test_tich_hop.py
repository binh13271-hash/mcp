"""Test đoạn ghép tich_hop/them_vao_server.py trong môi trường giả lập server (Codex R10).

Giả lập các hàm hạ tầng của server.py; KHÔNG thay việc chạy thật trên máy (Bước 6 skill chuyen-vien-mcp).
"""
import json
import pathlib

import openpyxl
import pytest

DA_DANG_KY, DA_DONG, DA_HYDRATE = [], [], []


def _moi_truong():
    def deco(fn):
        DA_DANG_KY.append(fn.__name__)
        return fn

    def nap(dd, **kw):
        wb = openpyxl.load_workbook(dd, **kw)
        goc = wb.close
        wb.close = lambda: (DA_DONG.append(kw.get("data_only", False)), goc())
        return wb

    def o_co_that(ws, d1=1, d2=None):
        # Giả lập _o_co_that() của server (generator ô thật theo dòng, cột) — bản ghép v7.3 dùng lại nó (G6).
        o = ws._cells
        for k in sorted(k for k in o if k[0] >= d1 and (d2 is None or k[0] <= d2)):
            yield o[k]

    import os
    g = {"_DECO_DOC": deco, "secure_path": lambda p: p, "ensure_file_hydrated": DA_HYDRATE.append,
         "_nap_workbook": nap, "_o_co_that": o_co_that, "json": json, "os": os, "__name__": "server_gia"}
    code = pathlib.Path(__file__).resolve().parents[1].joinpath("tich_hop", "them_vao_server.py").read_text(encoding="utf-8")
    exec(compile(code, "them_vao_server.py", "exec"), g)
    return g


@pytest.fixture
def sv():
    DA_DANG_KY.clear(), DA_DONG.clear(), DA_HYDRATE.clear()
    return _moi_truong()


@pytest.fixture
def file_kl(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KLCT"
    for o, v in {"B6": "TT", "C6": "HẠNG MỤC", "E7": "DÀI", "F7": "RỘNG", "I7": "G/ NHAU",
                 "J7": "RIÊNG", "K7": "CHUNG", "B10": 1, "C10": "+ Trừ hố 1.0x2.0 (10 hố)",
                 "E10": 1.0, "F10": 2.0, "I10": 10, "J10": "=-E10*F10"}.items():
        ws[o] = v
    wb.create_sheet("xxxxxxxx")
    p = tmp_path / "kl.xlsx"
    wb.save(p)
    return str(p)


def test_dang_ky_du_3_tool(sv):
    assert DA_DANG_KY == ["excel_quet_loi_khoi_luong", "cad_dem_doi_tuong", "cad_hatch_giao_doi_tuong"]


def test_quet_va_dong_ca_hai_workbook(sv, file_kl):
    ra = sv["excel_quet_loi_khoi_luong"](file_kl)
    assert "HE_SO_BO_SOT|KLCT!J10" in ra and "SHEET_LA|xxxxxxxx" in ra
    assert sorted(DA_DONG) == [False, True] and DA_HYDRATE == [file_kl]


def test_dong_workbook_ca_khi_loi_sheet(sv, file_kl):
    assert "Không có sheet" in sv["excel_quet_loi_khoi_luong"](file_kl, ten_sheet="KHONGCO")
    assert sorted(DA_DONG) == [False, True]


@pytest.mark.parametrize("cm,chu", [("{sai", "không phải JSON"), ("[1,2]", "phải là object"),
                                    (json.dumps({"kl": "J"}), "loại không biết"), (json.dumps({"kl_rieng": "j1"}), "cột không hợp lệ")])
def test_cot_map_sai_bao_ro(sv, file_kl, cm, chu):
    assert chu in sv["excel_quet_loi_khoi_luong"](file_kl, cot_map_json=cm)


def test_cad_chi_nhan_dxf_va_hydrate(sv, tmp_path):
    assert "Chỉ đọc .dxf" in sv["cad_dem_doi_tuong"](str(tmp_path / "a.dwg"))
    import ezdxf
    d = ezdxf.new()
    d.header["$INSUNITS"] = 6
    d.modelspace().add_line((0, 0), (3, 0))
    p = str(tmp_path / "a.dxf")
    d.saveas(p)
    assert "L|0|3.00" in sv["cad_dem_doi_tuong"](p) and DA_HYDRATE == [p]
    assert "ngoài miền" in sv["cad_hatch_giao_doi_tuong"](p, "LAT", "X", dung_sai=0.9)
